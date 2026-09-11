import json, subprocess, pathlib, datetime, hashlib
import openpyxl

root = pathlib.Path(__file__).resolve().parents[1]
source = pathlib.Path('C:/Users/Alejandro/Downloads/AUSENTISMO 2026 POR MES CALENDARIO V2.xlsx')
values = openpyxl.load_workbook(source, data_only=True, read_only=True)
formulas = openpyxl.load_workbook(source, data_only=False, read_only=True)
sheet = values['Agosto']
grid = list(sheet.iter_rows(max_row=37, max_col=20, values_only=True))
formula_grid = list(formulas['Agosto'].iter_rows(max_row=37,max_col=20,values_only=True))
print('Excel leido. Consultando el informe completo Buk y GeoVictoria...', flush=True)
proc = subprocess.run(['C:/xampp/php/php.exe', str(root/'tests/fte_compare_worker.php')], capture_output=True, text=True, encoding='utf-8', cwd=root, timeout=1800)
if proc.returncode:
    raise RuntimeError('No se obtuvo el informe completo. Codigo de salida: '+str(proc.returncode))
report = json.loads(proc.stdout)
portal = {r['cost_center_code']:r for r in report['cost_centers']}
metrics = [('Dotacion',5,'headcount'),('Horas teoricas',6,'theoretical_headcount_hours'),('Extra',7,'authorized_overtime_hours'),('Vacaciones',8,'vacation_hours'),('Ingresos/salidas',9,'employment_movement_hours'),('Licencias',10,'medical_leave_hours'),('Accidentes',11,'accident_hours'),('Permiso dia',12,'day_permission_hours'),('Permiso hora',13,'hour_permission_hours'),('Fallas/atrasos',14,'failure_delay_hours'),('Horas perdidas',15,'lost_hours'),('FTE',19,'fte')]
pending={'accident_hours','day_permission_hours','hour_permission_hours','authorized_overtime_hours','failure_delay_hours'}
def number(x): return float(x) if isinstance(x,(int,float)) else 0.0
def pv(row,key): return row.get(key,row.get('loss_components',{}).get(key))
def fmt(x): return 'No disponible' if x is None else f'{x:,.4f}'.rstrip('0').rstrip('.')
lines=['# Comparacion Excel RR.HH. y PortalGP: agosto 2026','',f'Consulta: {datetime.datetime.now().isoformat(timespec="seconds")}',f'Excel: {source.name}',f'SHA256: {hashlib.sha256(source.read_bytes()).hexdigest()}','',
'Valores Excel guardados y formulas inspeccionadas; no se modifico ni recalculo el original. Consulta Buk actual. GeoVictoria se interrumpio sin resultados recuperables: extras y atrasos NO estan comparados. El FTE del portal es parcial. Las APIs actuales pueden diferir del cierre historico.','',
'## Cobertura','',f"GeoVictoria: {report['attendance_diagnostics']}",'']
lines += ['- '+w for w in report['warnings']]
lines += ['', 'Los accidentes y permisos no tienen fuente automatica: sus ceros internos no se consideran coincidencias. Fallas/atrasos del portal representa solo atrasos y no tiene el mismo alcance completo que el Excel.', '', '## Totales', '', '|Concepto|Excel|Portal|Diferencia Portal - Excel|','|---|---:|---:|---:|']
for label,col,key in metrics:
    ev = number(grid[35][col]) if key!='fte' else number(grid[36][19])
    val = None if key in pending else pv(report['totals'],key)
    lines.append(f'|{label}|{fmt(ev)}|{fmt(val)}|{fmt(None if val is None else val-ev)}|')
matched=[]; missing=[]
for ri in range(8,35):
    row=grid[ri]; code=str(row[1] or '').strip(); name=str(row[2] or '').strip()
    if not name: continue
    pr=portal.get(code)
    lines += ['',f'## {name} ({code or "SIN CECO EN EXCEL"}) — fila {ri+1}','']
    if pr is None:
        missing.append((code,name)); lines += ['Sin correspondencia exacta por CECO. No se presume una equivalencia por nombre.']; continue
    matched.append(code)
    lines += ['|Concepto|Excel|Portal|Diferencia|','|---|---:|---:|---:|']
    for label,col,key in metrics:
        ev=number(row[col]); val=None if key in pending else pv(pr,key)
        lines.append(f'|{label}|{fmt(ev)}|{fmt(val)}|{fmt(None if val is None else val-ev)}|')
lines += ['', '## Areas sin correspondencia', '', 'Excel: '+repr(missing),'Portal: '+repr([(c,r['cost_center_name']) for c,r in portal.items() if c not in matched]),'', '## Formulas y cuadratura del Excel','']
for ri in range(8,35):
    row=grid[ri]; expected=sum(number(x) for x in row[8:15]); actual=number(row[15]); calculated=(number(row[6])+number(row[7])-expected)/number(grid[2][8])
    if abs(actual-expected)>0.00001 or abs(number(row[19])-calculated)>0.00001:
        lines.append(f'- Fila {ri+1}: perdidas {actual} frente a suma {expected}; FTE {row[19]} frente a {calculated}.')
lines += [f'- R36 contiene `{formula_grid[35][17]}` y vale {grid[35][17]}. El indice global calculado desde horas es {1-number(grid[35][15])/number(grid[35][6]):.8%}.', '', '## Formulas de referencia', '', '```']
for ri in [8,11,16,35,36]:
    lines.append(str(ri+1)+': '+repr({openpyxl.utils.get_column_letter(ci+1):v for ci,v in enumerate(formula_grid[ri]) if isinstance(v,str) and v.startswith('=')}))
lines += ['```']
out=root/'rrhh/fte/FTE_COMPARACION_EXCEL_AGOSTO_2026.md'
out.write_text('\n'.join(lines)+'\n',encoding='utf-8')
print('Informe: '+str(out),flush=True)
print(json.dumps({'matched':len(matched),'missing_excel':missing,'portal_only':[(c,r['cost_center_name']) for c,r in portal.items() if c not in matched],'totals':report['totals'],'diagnostics':report['attendance_diagnostics'],'warnings':report['warnings']},ensure_ascii=True),flush=True)
